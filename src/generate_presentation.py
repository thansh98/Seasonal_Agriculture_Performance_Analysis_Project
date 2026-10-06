"""
PowerPoint Automation Module
Seasonal Agriculture Performance Analysis
AICTE | IBM SkillsBuild Data Analytics with AI Internship Program 2026
Implementation Partner: BharatCares | Trainer: Mr. Kartik Hooda
Fills in Major Project PPT Template with content, results, metrics, and high-res figures.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation(template_path: str, figures_dir: str, output_path: str):
    prs = Presentation(template_path)
    print(f"[INFO] Loaded template with {len(prs.slides)} slides.")

    # Colors
    DARK_BLUE = RGBColor(16, 44, 87)
    GREEN = RGBColor(43, 138, 62)
    DARK_GRAY = RGBColor(50, 50, 50)
    LIGHT_GRAY = RGBColor(240, 240, 240)

    # Slide 1: Title Slide
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.name == 'Title 3' or 'Title' in shape.name:
            shape.text_frame.text = "Seasonal Agriculture Performance Analysis\nAICTE | IBM SkillsBuild Internship 2026"
            for p in shape.text_frame.paragraphs:
                p.font.size = Pt(26)
                p.font.bold = True
                p.font.color.rgb = DARK_BLUE
        elif 'Placeholder 1' in shape.name:
            shape.text_frame.text = "Student Name: [Your Name]\nCollege: [Your College / University Name]\nPartner: BharatCares | Trainer: Mr. Kartik Hooda"
            for p in shape.text_frame.paragraphs:
                p.font.size = Pt(14)
                p.font.color.rgb = DARK_GRAY
        elif 'TextBox 2' in shape.name:
            shape.text_frame.text = "AICTE STU ID: [Your AICTE ID] | IBM SkillsBuild ID: [Your ID]\nTopic: Data Analytics with AI: Foundation to Implementation"
            for p in shape.text_frame.paragraphs:
                p.font.size = Pt(13)
                p.font.color.rgb = DARK_GRAY

    # Slide 2: Problem Statement
    s2 = prs.slides[1]
    # Keep problem statement title and text crisp and readable
    for shape in s2.shapes:
        if shape.has_text_frame and shape.name != 'Title 3':
            tf = shape.text_frame
            tf.clear()
            p1 = tf.paragraphs[0]
            p1.text = "Agricultural activities are profoundly influenced by seasonal variations in climate, resource availability, and soil dynamics:"
            p1.font.size = Pt(15)
            p1.font.bold = True
            
            bullets = [
                "Information Gap: Raw agricultural records do not explain how production, input costs, and net farm incomes fluctuate across seasons.",
                "Resource Disparity: Water and fertilizer usage vary widely between Kharif (monsoon), Rabi (winter), and Zaid (summer), impacting sustainability.",
                "Economic Volatility: High summer input costs often turn Zaid crops into loss-making ventures despite adequate market prices.",
                "Core Challenge: To comprehensively investigate the 4,000-farm multi-state agricultural dataset, quantify seasonal performance variations using rigorous statistical tests (ANOVA, Kruskal-Wallis), and build high-precision predictive ML models (R2 > 97%) for evidence-based agricultural planning."
            ]
            for b in bullets:
                p = tf.add_paragraph()
                p.text = "• " + b
                p.font.size = Pt(13)
                p.font.color.rgb = DARK_GRAY

    # Slide 3: Project Description
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if 'Title' in shape.name:
            shape.text_frame.text = "Project Description & Methodology"
            shape.text_frame.paragraphs[0].font.size = Pt(24)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Comprehensive End-to-End Data Analytics & Machine Learning Pipeline:"
            p0.font.size = Pt(15)
            p0.font.bold = True
            sections = [
                "Dataset Scope: 4,000 farm records spanning 8 major Indian states, 8 crops (Cereals, Cash Crops, Pulses), and 3 distinct seasons (Kharif, Rabi, Zaid) across 28 parameters.",
                "Robust Data Engineering: Solved missing values in Rainfall (State-Season median), Soil Moisture (Crop-Season median), and Yield (reconstructed via Production / Farm Area). Engineered ROI, Profit Margin %, Cost/Ha, and Water Productivity.",
                "Statistical Rigor: Conducted One-Way ANOVA, Kruskal-Wallis (Yield H=70.59, p<1e-15; Profit H=101.93, p<1e-22), Chi-Square tests, and pairwise Mann-Whitney U tests with Bonferroni correction.",
                "Predictive AI Modeling: Trained and benchmarked Random Forest, Gradient Boosting, and Ridge Regression models, achieving 97.5% R2 in Yield prediction and 98.1% R2 in Farm Profitability prediction.",
                "Interactive Web Application: Designed a Streamlit decision-support dashboard allowing farmers and planners to simulate seasonal interventions and irrigation shifts in real time."
            ]
            for sec in sections:
                p = tf.add_paragraph()
                p.text = "• " + sec
                p.font.size = Pt(12)
                p.font.color.rgb = DARK_GRAY

    # Slide 4: End Users
    s4 = prs.slides[3]
    for shape in s4.shapes:
        if shape.has_text_frame and shape.name != 'Title 3':
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Target Stakeholders & Decision-Making Beneficiaries:"
            p0.font.size = Pt(15)
            p0.font.bold = True
            users = [
                "Individual Farmers & FPOs: Optimize crop-season selection, avoid unprofitable summer ventures, and adopt high-efficiency Drip irrigation to boost margins by up to 2.3x.",
                "Ministry of Agriculture & Policy Planners: Formulate targeted seasonal fertilizer subsidies, optimize regional canal water allocations, and deploy MSP support for staple grains.",
                "Agritech & Micro-Irrigation Companies: Identify high-impact conversion clusters (e.g. converting flood-irrigated zones to drip in water-stressed districts).",
                "Agricultural Financial & Insurance Institutions: Accurately price seasonal crop insurance based on validated pest risk models (monsoon pest probability = 54.5%) and credit underwriting.",
                "Agronomy Researchers & Extension Workers: Utilize empirical seasonal benchmarks to design location-specific farming practice packages."
            ]
            for u in users:
                p = tf.add_paragraph()
                p.text = "• " + u
                p.font.size = Pt(13)
                p.font.color.rgb = DARK_GRAY

    # Slide 5: Technology Used
    s5 = prs.slides[4]
    for shape in s5.shapes:
        if shape.has_text_frame and 'Title' not in shape.name:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Modern Data Science & Analytics Tech Stack:"
            p0.font.size = Pt(15)
            p0.font.bold = True
            stack = [
                "Programming Language: Python 3.12 (High performance data structures & vectorized execution)",
                "Data Manipulation & Engineering: Pandas 2.2, NumPy (Data cleaning, outlier analysis, per-hectare metrics)",
                "Statistical Inference: SciPy Stats (One-Way ANOVA, Kruskal-Wallis H, Mann-Whitney U, Chi-Square)",
                "Machine Learning: Scikit-Learn (Random Forest, Gradient Boosting Regressor, Ridge Regression, K-Fold CV)",
                "Data Visualization: Matplotlib, Seaborn (Publication-ready 300 DPI figures, multi-panel infographics)",
                "Interactive Dashboard: Streamlit 1.38 (Interactive reactive UI for predictive scenario simulation)",
                "Interactive Documentation: Jupyter Notebook (.ipynb), NBConvert, Git/GitHub for version control"
            ]
            for s in stack:
                p = tf.add_paragraph()
                p.text = "• " + s
                p.font.size = Pt(13)
                p.font.color.rgb = DARK_GRAY

    # Slide 6: Results Overview
    s6 = prs.slides[5]
    for shape in s6.shapes:
        if shape.has_text_frame and 'Title' not in shape.name:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Summary of Key Empirical Findings Across 4,000 Farms:"
            p0.font.size = Pt(14)
            p0.font.bold = True
            findings = [
                "Seasonal Profit Divergence: Kharif delivers the highest mean profit (INR 178,914), followed by Rabi (INR 87,689). In contrast, Zaid results in an average net loss of -INR 24,805 due to high irrigation water demand and depressed yields.",
                "Yield Variations: Average yields drop progressively from Kharif (5.63 t/ha) to Rabi (5.09 t/ha) and Zaid (4.63 t/ha) (Kruskal-Wallis H = 70.59, p = 4.69e-16).",
                "Water Productivity Premium: Drip irrigation achieves 7.26 t/ha yield in Kharif and maintains profitability in Zaid (+INR 21,291), whereas Flood irrigation causes heavy Zaid losses (-INR 69,786).",
                "Disease & Pest Dynamics: High humidity (71.8%) and rainfall (852 mm) in Kharif drive pest risk to 54.5%, requiring 35% higher preventive management than winter (40.5%).",
                "Machine Learning Precision: Gradient Boosting models achieved 97.5% R2 for Yield prediction and 98.1% R2 for Farm Profitability."
            ]
            for f in findings:
                p = tf.add_paragraph()
                p.text = "• " + f
                p.font.size = Pt(12)
                p.font.color.rgb = DARK_GRAY

    # Slide 7: Results - Climate & Yield
    s7 = prs.slides[6]
    for shape in s7.shapes:
        if shape.name == 'Title 3' or 'Title' in shape.name:
            shape.text_frame.text = "Results: Seasonal Climate & Crop Yield Variations"
            shape.text_frame.paragraphs[0].font.size = Pt(22)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Key Observations:\n• Rainfall drops 65% from Kharif (852mm) to Zaid (299mm).\n• Avg Temperature rises to 31.0°C in Zaid vs 23.5°C in Rabi.\n• Cash crops (Sugarcane, Chilli) deliver superior yields across all seasons."
            p0.font.size = Pt(11)
            p0.font.color.rgb = DARK_GRAY
    # Add image
    fig_path = os.path.join(figures_dir, 'seasonal_distributions.png')
    if os.path.exists(fig_path):
        s7.shapes.add_picture(fig_path, Inches(0.8), Inches(2.2), width=Inches(8.4))

    # Slide 8: Results - Irrigation & Resource Efficiency
    s8 = prs.slides[7]
    for shape in s8.shapes:
        if shape.name == 'Title 3' or 'Title' in shape.name:
            shape.text_frame.text = "Results: Water Efficiency & Irrigation Impact"
            shape.text_frame.paragraphs[0].font.size = Pt(22)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Key Observations:\n• Drip irrigation outperforms flood irrigation across all seasons with 2.3x higher profit.\n• Flood and Rainfed systems turn severely loss-making in Zaid (-INR 69k to -79k).\n• Drip and Sprinkler systems protect water efficiency (4.41 to 5.89 t/1000m3)."
            p0.font.size = Pt(11)
            p0.font.color.rgb = DARK_GRAY
    fig_path = os.path.join(figures_dir, 'irrigation_efficiency.png')
    if os.path.exists(fig_path):
        s8.shapes.add_picture(fig_path, Inches(0.8), Inches(2.2), width=Inches(8.4))

    # Slide 9: Results - Economic Performance
    s9 = prs.slides[8]
    for shape in s9.shapes:
        if shape.name == 'Title 3' or 'Title' in shape.name:
            shape.text_frame.text = "Results: Crop Profitability & Economic Disparities"
            shape.text_frame.paragraphs[0].font.size = Pt(22)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Key Observations:\n• Chilli and Sugarcane are highly profitable in all 3 seasons (>INR 440k net profit).\n• Grain crops (Wheat, Rice, Maize) face margin compression under rising cost per hectare.\n• Kharif has highest proportion of profitable farms (68%), whereas Zaid has 54% loss-making farms."
            p0.font.size = Pt(11)
            p0.font.color.rgb = DARK_GRAY
    fig_path = os.path.join(figures_dir, 'crop_profitability_matrix.png')
    if os.path.exists(fig_path):
        s9.shapes.add_picture(fig_path, Inches(1.5), Inches(2.2), width=Inches(7.0))

    # Slide 10: Results - ML Predictive Models
    s10 = prs.slides[9]
    for shape in s10.shapes:
        if shape.name == 'Title 3' or 'Title' in shape.name:
            shape.text_frame.text = "Results: Machine Learning Modeling & Key Drivers"
            shape.text_frame.paragraphs[0].font.size = Pt(22)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Key Observations:\n• Gradient Boosting achieved Yield R2 = 97.5% (MAE 0.64 t/ha) & Profit R2 = 98.1%.\n• Primary yield drivers: Crop type (Sugarcane/cash crops), Soil pH, Rainfall, Nitrogen, and Irrigation Method.\n• Primary profit drivers: Production volume, Market price per tonne, Soil pH, and Input cost containment."
            p0.font.size = Pt(11)
            p0.font.color.rgb = DARK_GRAY
    fig_path = os.path.join(figures_dir, 'feature_importance.png')
    if os.path.exists(fig_path):
        s10.shapes.add_picture(fig_path, Inches(1.5), Inches(2.2), width=Inches(7.0))

    # Slide 11: Future Scope
    s11 = prs.slides[10]
    for shape in s11.shapes:
        if shape.has_text_frame and 'Title' not in shape.name:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Strategic Roadmap & Future Scope of Expansion:"
            p0.font.size = Pt(15)
            p0.font.bold = True
            scopes = [
                "Remote Sensing Integration: Incorporate Sentinel-2 and Landsat multispectral imagery for real-time NDVI vegetation health and soil moisture indexing.",
                "IoT & Telemetry Integration: Connect low-cost soil moisture probes and weather station telemetry for automated precision irrigation scheduling.",
                "Deep Learning & Computer Vision: Train CNN models for real-time smartphone-based pest and disease leaf lesion classification during high-risk monsoon periods.",
                "Dynamic APMC Market Price Forecasting: Incorporate Prophet/LSTM time-series models to predict seasonal wholesale commodity prices and guide optimal harvest timing.",
                "Farmer-Centric Mobile PWA: Deploy multilingual offline-first progressive web application delivering hyper-local advisory to smallholder farmers."
            ]
            for sc in scopes:
                p = tf.add_paragraph()
                p.text = "• " + sc
                p.font.size = Pt(13)
                p.font.color.rgb = DARK_GRAY

    # Slide 12: GitHub Link
    s12 = prs.slides[11]
    # Keep github link clean and informative
    for shape in s12.shapes:
        if 'TextBox' in shape.name or (shape.has_text_frame and 'Title' not in shape.name):
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Project Source Code Repository:\nhttps://github.com/anshp/Seasonal-Agriculture-Performance-Analysis"
            p0.font.size = Pt(16)
            p0.font.bold = True
            p0.font.color.rgb = DARK_BLUE
            p1 = tf.add_paragraph()
            p1.text = "Includes: Complete Jupyter Notebook (.ipynb), Streamlit Web App (app.py), Preprocessing & Statistical Modules, ML Models, and High-Res Figures."
            p1.font.size = Pt(13)
            p1.font.color.rgb = DARK_GRAY

    # Slide 13: Course Completion Certificate Placeholder
    s13 = prs.slides[12]
    for shape in s13.shapes:
        if 'Title' in shape.name:
            shape.text_frame.text = "IBM SkillsBuild / BharatCares Completion Certificate"
            shape.text_frame.paragraphs[0].font.size = Pt(22)
            shape.text_frame.paragraphs[0].font.bold = True
            shape.text_frame.paragraphs[0].font.color.rgb = DARK_BLUE
        elif shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "[Insert your IBM SkillsBuild / BharatCares Data Analytics with AI Course Completion Certificate screenshot here]\nProgram: Data Analytics with AI: Foundation to Implementation\nTrainer: Mr. Kartik Hooda | Company: BharatCares"
            p0.font.size = Pt(15)
            p0.font.bold = True
            p0.font.color.rgb = DARK_GRAY

    # Slide 14: Thank You
    s14 = prs.slides[13]
    for shape in s14.shapes:
        if shape.has_text_frame and 'Title' not in shape.name:
            tf = shape.text_frame
            tf.clear()
            p0 = tf.paragraphs[0]
            p0.text = "Thank You!\nQuestions & Feedback Welcome\nAICTE | IBM SkillsBuild Data Analytics with AI Internship Program 2026\nImplementation Partner: BharatCares | Trainer: Mr. Kartik Hooda"
            p0.font.size = Pt(17)
            p0.font.bold = True
            p0.font.color.rgb = DARK_BLUE

    prs.save(output_path)
    print(f"[SUCCESS] Final Presentation generated and saved to: {output_path}")

def main():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    template_path = os.path.join(base_dir, 'data', 'Major_Project_PPT_Submission_Template.pptx')
    figures_dir = os.path.join(base_dir, 'reports', 'figures')
    output_path = os.path.join(base_dir, 'reports', 'IBM_SkillsBuild_Seasonal_Agriculture_Performance_Analysis.pptx')
    
    create_presentation(template_path, figures_dir, output_path)

if __name__ == '__main__':
    main()
